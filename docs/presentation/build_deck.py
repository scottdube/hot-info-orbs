"""Fill the 12 empty SmComp-Dark slides of base.pptx with the HOT Info Orbs intro."""
import sys, copy
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from PIL import Image

import os
D = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'assets') + '/'
src, out = sys.argv[1], sys.argv[2]

WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GOLD = RGBColor(0xFF, 0xC0, 0x00)
ICE = RGBColor(0xCA, 0xDC, 0xFC)
FRAME = RGBColor(0x5B, 0x9B, 0xD5)
GREY = RGBColor(0xBF, 0xC7, 0xD5)

prs = Presentation(src)
slides = list(prs.slides)
assert len(slides) == 12, len(slides)


def pic(slide, path, x, y, w=None, h=None, frame=True, fit=None):
    """Add a picture; fit=(w,h) scales it to fit inside that box, centered."""
    im = Image.open(path)
    ar = im.width / im.height
    if fit:
        fw, fh = fit
        if fw / fh > ar:
            h = fh; w = fh * ar
        else:
            w = fw; h = fw / ar
        x = x + (fw - w) / 2
        y = y + (fh - h) / 2
    elif w is None:
        w = h * ar
    elif h is None:
        h = w / ar
    p = slide.shapes.add_picture(path, Inches(x), Inches(y), Inches(w), Inches(h))
    if frame:
        p.line.color.rgb = FRAME
        p.line.width = Pt(1)
    return p, (x, y, w, h)


def tb(slide, x, y, w, h, paras, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, name=None):
    """paras: list of (text, size, bold, color) or list of runs [(text,size,bold,color),...]
    with an optional dict as the last element for space_after."""
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        box.name = name
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.05)
    tf.margin_top = tf.margin_bottom = Inches(0.03)
    tf.vertical_anchor = anchor
    first = True
    for para in paras:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = align
        runs = para if isinstance(para, list) else [para]
        opts = {}
        if runs and isinstance(runs[-1], dict):
            opts = runs[-1]; runs = runs[:-1]
        for text, size, bold, color in runs:
            r = p.add_run()
            r.text = text
            r.font.name = 'Arial'
            r.font.size = Pt(size)
            r.font.bold = bold
            r.font.color.rgb = color
        p.space_after = Pt(opts.get('after', 4))
        p.space_before = Pt(opts.get('before', 0))
    return box


def title(slide, text, sub=None):
    slide.shapes.clone_layout_placeholders(slide.slide_layout)
    for ph in slide.placeholders:
        t = ph.placeholder_format.type
        if str(t).startswith('TITLE') or str(t).startswith('CENTER_TITLE'):
            ph.text_frame.text = text
        elif str(t).startswith('SUBTITLE'):
            if sub is None:
                ph._element.getparent().remove(ph._element)
            else:
                ph.text_frame.text = sub
    return slide


def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


def caption(slide, x, y, w, head, body=None, size=14, h=0.9):
    paras = [(head, size, True, GOLD)]
    if body:
        paras.append((body, size - 1, False, WHITE))
    return tb(slide, x, y, w, h, paras)


# ---------- 1 Title ----------
s = slides[0]
title(s, 'HOT Info Orbs', 'Five round screens that show what you want to know')
for ph in s.placeholders:
    t = str(ph.placeholder_format.type)
    if t.startswith('CENTER_TITLE'):
        ph.left, ph.top, ph.width, ph.height = Inches(0.6), Inches(0.35), Inches(11.2), Inches(1.4)
        for p in ph.text_frame.paragraphs: p.alignment = PP_ALIGN.LEFT
    elif t.startswith('SUBTITLE'):
        ph.left, ph.top, ph.width, ph.height = Inches(0.6), Inches(1.75), Inches(11.2), Inches(0.8)
        for p in ph.text_frame.paragraphs:
            p.alignment = PP_ALIGN.LEFT
            for r in p.runs: r.font.size = Pt(24)
pic(s, D + 'c-dim-before.jpg', 0.6, 2.7, w=12.1)
tb(s, 0.6, 6.3, 12.1, 0.9, [
    [('Hands-On Tech, Small Computers group', 18, True, GOLD), ('   ·   October 27, 2026   ·   Scott Dube', 18, False, WHITE)],
    ('A club build: a board designed by Yang and Scott, firmware and docs on GitHub', 16, False, GREY),
])
notes(s, "Open on the orb itself, running on the table. "
      "This is the HOT Info Orbs: five round screens on one board, run by a Wi-Fi dev board about the size of a postage stamp. "
      "They show the time, the weather and the markets, and they keep themselves up to date. "
      "It's a Hands-On Tech club build: anyone in the club can make one and set it up from a phone or computer. "
      "Tonight: what it is, where it came from, a tour of the pages, how it's built, and how the club build will run.")

# ---------- 2 What it is ----------
s = slides[1]
title(s, 'What it is')
pic(s, D + 'hero-4k-band.jpg', 0.6, 1.5, w=12.1)
cols = [
    ('Five round screens', 'Color displays an inch and a quarter across, on one board. Each page uses all five.'),
    ('One small computer', 'A XIAO ESP32-S3 with Wi-Fi fetches the time, the weather, the markets and your own data.'),
    ('Nothing to install', 'A USB cable for power, three buttons, and a settings page you open on your phone.'),
]
for i, (h, b) in enumerate(cols):
    caption(s, 0.6 + i * 4.1, 4.85, 3.9, h, b, size=18, h=1.8)
notes(s, "Five round screens on one board, run by a Wi-Fi dev board about the size of a postage stamp. "
      "They show the time, the weather and the markets, and they keep themselves up to date. "
      "There is nothing to install on a computer: it powers from a USB cable, and every setting lives on a web page the orb serves itself.")

# ---------- 3 Where it came from ----------
s = slides[2]
title(s, 'Where it came from')
items = [
    (D + 'brett-43.jpg', 'The original: Info Orbs', "Brett Turner designed the Info Orbs and shared the plans and software for free.", 'Photo: Brett Turner, brett.tech'),
    (D + 'S06-hackerboxes-0129-product-image.png', 'The kit: HackerBox 0129 “Orbital”', "HackerBoxes, a monthly electronics kit, sent out its own version.", 'Image: HackerBoxes'),
    (D + 'S07-donald-champion-hackerbox-0129-IMG_2820.jpg', 'Built in the club', "A couple of our members built that kit. We liked it enough to make our own: a new board, new features, a build anyone can follow.", 'Photo: Donald Champion'),
]
for i, (path, h, b, cr) in enumerate(items):
    x = 0.6 + i * 4.15
    pic(s, path, x, 1.6, fit=(3.9, 2.95))
    tb(s, x, 4.65, 3.9, 2.2, [(h, 18, True, GOLD), (b, 15, False, WHITE), (cr, 12, False, GREY, )])
notes(s, "The idea isn't ours. Brett Turner designed the original Info Orbs and shared the plans and software for free. "
      "Later HackerBoxes, a monthly electronics kit, sent out a version of it called Orbital. "
      "A couple of our members built that kit; these are Donald Champion's photos of his. "
      "We liked it enough to make our own: a new board, new features, and a build anyone in the club can follow.")

# ---------- 4 Pages: clock and weather ----------
s = slides[3]
title(s, 'The pages: clock and weather')
rows = [
    (D + 'c-clock-standard.jpg', 'Clock, LED digits'),
    (D + 'c-clock-nixie.jpg', 'Clock, Nixie face (middle button)'),
    (D + 'c-dim-before.jpg', 'Weather: now, today, and the next three days'),
]
y = 1.55
for path, cap in rows:
    _, (x, yy, w, h) = pic(s, path, 0.6, y, w=5.1)
    tb(s, 0.6, y + h + 0.02, 5.1, 0.33, [(cap, 14, True, GOLD)])
    y += h + 0.41
_, (x, yy, w, h) = pic(s, D + 'sunmoon-close.jpg', 6.1, 1.55, w=6.75)
tb(s, 6.1, yy + h + 0.15, 6.75, 1.7, [
    ('Sunrise, sunset and tonight’s moon', 14, True, GOLD),
    ('The clock spreads across all five orbs: 12- or 24-hour, in any colors you like.', 15, False, WHITE),
    ('Weather for any town, or straight from your own Tempest station. Up to two stations.', 15, False, WHITE),
])
notes(s, "The clock spreads across all five orbs. Press the middle button for a nixie tube face. "
      "Twelve or twenty-four hour, in any colors you like, and you can build in a custom face of your own. "
      "The weather page covers any town: what it's doing now and the next three days. "
      "If you own a Tempest weather station, it can drive the display instead; you can add up to two. "
      "Sunrise and sunset sit above the weather icon, and tonight's moon is drawn below it.")

# ---------- 5 Pages: stocks and your own data ----------
s = slides[4]
title(s, 'The pages: stocks and your own data')
rows = [
    (D + 'c-stocks-open.jpg', 'Stocks', 'Five tickers, one per orb: price, change and the day’s range. Stocks, funds or crypto.'),
    (D + 'c-wide-mqtt.jpg', 'MQTT', 'Five live values from Home Assistant, or anything else that speaks MQTT.'),
    (D + 'c-wide-webdata.jpg', 'Web Data', 'A small JSON file from any web address, laid out the way you want.'),
]
y = 1.55
for path, h, b in rows:
    _, (x, yy, w, hh) = pic(s, path, 0.6, y, w=4.6)
    tb(s, 5.5, y, 7.35, hh, [(h, 18, True, GOLD), (b, 15, False, WHITE)], anchor=MSO_ANCHOR.MIDDLE)
    y += hh + 0.18
tb(s, 0.6, y - 0.05, 12.25, 0.5, [('Left and right buttons step through the pages. Leave it alone and it rotates on its own. The last two pages are ones you build yourself.', 14, False, GREY)])
notes(s, "Five tickers, one per orb, with the price, the change and the day's range: stocks, funds or crypto. "
      "Left and right step through the pages; leave it alone and it rotates on its own. "
      "And for coders, two pages you build yourself. The MQTT page shows live values from Home Assistant, or anything else using MQTT. "
      "The Web Data page pulls a small JSON file from any web address.")

# ---------- 6 How it's built ----------
s = slides[5]
title(s, 'How it’s built')
pic(s, D + 'b24-front-crop.png', 1.15, 1.45, w=11.0, frame=False)
calls = [
    ('XIAO ESP32-S3 or S3 Plus', 'the brain, Wi-Fi built in, in a socket'),
    ('Five 1.28 in displays', 'GC9A01 round color TFTs, socketed'),
    ('Three buttons', 'pages, clock face, wake'),
    ('USB-C', 'power and the first firmware load'),
    ('Under 2 W', '1.7 W measured at 5 V'),
]
for i, (h, b) in enumerate(calls):
    caption(s, 0.6 + i * 2.45, 4.75, 2.35, h, b, size=16, h=1.5)
tb(s, 0.6, 6.3, 12.1, 0.5, [[('Board by Yang and Scott', 16, True, GOLD), ('   ·   made by JLCPCB   ·   every file on GitHub', 16, False, WHITE)]])
notes(s, "Here's what's inside. The brain is a XIAO ESP32-S3 or S3 Plus, with Wi-Fi built in. It sits in a socket, so it lifts out. "
      "Five round color displays, an inch and a quarter across, share one board; three buttons along the edge. "
      "Yang and I designed the board that ties it together, and JLCPCB makes it. "
      "It runs off a USB cable and draws under two watts. The board files and the software are all open, on GitHub.")

# ---------- 7 Why the displays run on 5 V ----------
s = slides[6]
title(s, 'Why the displays run on 5 V')
pic(s, D + 'module-2026-10-17.jpg', 0.6, 1.55, h=4.5)
pic(s, D + 'display-back-2026-10-17.jpg', 4.2, 1.55, h=4.5)
tb(s, 0.6, 6.1, 3.4, 0.6, [('The ESP32 module, hottest point', 13, True, GOLD)])
tb(s, 4.2, 6.1, 3.4, 0.6, [('The backs of three displays', 13, True, GOLD)])
stats = [
    ('✔  43.6 °C', 'our module, after hours of running', GOLD),
    ('✔  37.9 °C', 'our display backs', GOLD),
    ('✖  82 °C', 'the kit’s module: “too hot to touch”, a builder’s report', GREY),
]
for i, (big, small, col) in enumerate(stats):
    tb(s, 8.0, 1.5 + i * 1.3, 4.85, 1.25, [(big, 36, True, col), (small, 14, False, WHITE)])
tb(s, 8.0, 5.45, 4.85, 1.3, [('On the kit every display ran off the ESP’s regulator. Here each display has its own 5 V feed, so they carry their own share and stay warm, not hot.', 15, False, WHITE)])
notes(s, "The displays run on five volts, each with its own regulator. On the original kit they all ran off the ESP's regulator, "
      "and builders found it too hot to touch: one report put the module at 82 degrees C running only a graphics demo. "
      "On ours the thermal camera reads 43.6 at the module's hottest point after hours of running, and 37.9 on the display backs. "
      "Surface readings, room at 75 F. The displays carry their own share now, and stay warm, not hot.")

# ---------- 8 Setting it up ----------
s = slides[7]
title(s, 'Setting it up from your phone')
pic(s, D + 'panel-rotation-clock.png', 0.6, 1.55, h=5.2)
pic(s, D + 'panel-dim-off.png', 4.5, 1.55, h=3.95)
tb(s, 0.6, 6.78, 3.7, 0.3, [('Pages and clock', 12, True, GOLD)])
tb(s, 4.5, 5.52, 3.7, 0.3, [('Dim and screens-off times', 12, True, GOLD)])
_, (_, _, _, sh) = pic(s, D + 'panel-startup.png', 8.45, 1.55, w=4.4)
tb(s, 8.45, 1.55 + sh + 0.04, 4.4, 0.3, [('Your own start-up picture', 12, True, GOLD)])
tb(s, 8.45, 1.55 + sh + 0.45, 4.4, 2.3, [
    ('Tick the pages you want; untick the rest', 14, False, WHITE),
    ('Your town, tickers, time zone, colors', 14, False, WHITE),
    ('Dims in the evening, screens off at bedtime; any button wakes it', 14, False, WHITE),
    ('Firmware loads once over USB, then updates over Wi-Fi from this page', 14, False, WHITE),
])
notes(s, "You load the software once, over USB. After that there's no cable. "
      "Settings live on a web page you open from your phone or computer: your town, your tickers, time zone, colors. Save, and the orb restarts with them. "
      "And you pick the pages: not interested in the markets? Untick Stocks and it drops out of the rotation. Tick just one page and that page stays put. "
      "Software updates come over Wi-Fi from the same page; if one goes wrong the orb falls back to the version that worked. "
      "It dims in the evening, can switch its screens off at bedtime, and any button wakes it. And it starts up with any picture you like.")

# ---------- 9 What's next ----------
s = slides[8]
title(s, 'What’s next')
pic(s, D + 'b24-top2-crop.png', 0.6, 1.5, w=7.4, frame=False)
tb(s, 8.3, 1.6, 4.55, 2.2, [
    ('✔  Board rev 2.4: design locked', 20, True, GOLD),
    ('XIAO ESP32-S3 or S3 Plus, 46 mm display pitch to match the enclosure, buttons on top. This is the board the club builds.', 15, False, WHITE),
])
pic(s, D + 'enclosure-concept.png', 0.6, 4.1, w=7.4)
tb(s, 8.3, 4.2, 4.55, 2.5, [
    ('○  Enclosure: concept, by Yang', 20, True, GOLD),
    ('“Winged Speedster”: brass rings around each screen, finned end caps, all 3D printed. Still a design.', 15, False, WHITE),
])
notes(s, "Two things. The board has just moved to revision 2.4, and the design is locked: a XIAO ESP32-S3 or S3 Plus, the five displays at 46 millimeters to match the enclosure, buttons on top. "
      "And a case: Yang's Winged Speedster concept, brass rings around each screen and finned end caps, all 3D printed. That one is still a design.")

# ---------- 10 The club build ----------
s = slides[9]
title(s, 'The club build')
steps = [
    ('1  Get the parts', 'Board, XIAO ESP32-S3 or S3 Plus, five displays, sockets, buttons, a USB-C cable.'),
    ('2  Build night', 'Sockets and headers only: through-hole soldering, one evening.'),
    ('3  Load the firmware', 'Once, over USB. After that, updates from your phone.'),
    ('4  Set it up', 'Pick your pages on the settings page. No code.'),
]
for i, (h, b) in enumerate(steps):
    tb(s, 0.6, 1.6 + i * 1.15, 6.9, 1.1, [(h, 20, True, GOLD), (b, 15, False, WHITE)])
tb(s, 0.6, 6.25, 6.9, 0.5, [('Starting in January. Say tonight if you want in.', 18, True, WHITE)])
_, (x, yy, w, h) = pic(s, D + 'boot-close-hand.jpg', 7.85, 1.6, w=5.0)
tb(s, 7.85, yy + h + 0.15, 5.0, 1.6, [('Boots with the club logo', 14, True, GOLD), ('or any picture you upload. The module lifts out of its socket, so nothing is soldered that can’t be swapped.', 14, False, WHITE)])
notes(s, "Here is how the club build runs. You get the parts: the board, a XIAO ESP32-S3 or S3 Plus, five displays, sockets, buttons and a cable. "
      "Build night is sockets and headers only, through-hole soldering, one evening. "
      "Then the firmware loads once over USB; from then on updates come from your phone. Set it up on the settings page, no code. "
      "We start in January. Say tonight if you want in, so we know how many kits to put together.")

# ---------- 11 Links and credits ----------
s = slides[10]
title(s, 'Links and credits')
pic(s, D + 'qr-github.png', 0.6, 1.55, w=3.0, frame=False)
tb(s, 3.8, 1.55, 4.0, 3.0, [
    ('Everything is on GitHub', 18, True, GOLD),
    ('github.com/scottdube/hot-info-orbs', 16, False, WHITE),
    ('Board files, firmware, build guide', 14, False, GREY, ),
    ('', 8, False, WHITE),
    ('Feature sheet', 18, True, GOLD),
    ('scottdube.github.io/hot-info-orbs/features.html', 14, False, WHITE),
])
tb(s, 8.3, 1.55, 4.55, 5.0, [
    ('Thanks to', 18, True, GOLD),
    [('Brett Turner', 15, True, WHITE), ('  Info Orbs, the original', 15, False, WHITE)],
    [('HackerBoxes', 15, True, WHITE), ('  #0129 Orbital', 15, False, WHITE)],
    [('Donald Champion', 15, True, WHITE), ('  HackerBox build photos', 15, False, WHITE)],
    [('Yang and Scott', 15, True, WHITE), ('  board design', 15, False, WHITE)],
    [('Yang', 15, True, WHITE), ('  enclosure', 15, False, WHITE)],
    [('The members who built the first kit', 15, True, WHITE)],
])
notes(s, "Everything's on GitHub, at the link on screen: board files, firmware and the build guide. The feature sheet lists every page and setting. "
      "Thanks to Brett Turner, to HackerBoxes, to Donald Champion for the build photos, and to the members who built the first kit.")

# ---------- 12 End ----------
s = slides[11]
title(s, 'Questions?', 'The orb is on the table. Press the buttons.')
for ph in s.placeholders:
    t = str(ph.placeholder_format.type)
    if t.startswith('CENTER_TITLE'):
        ph.left, ph.top, ph.width, ph.height = Inches(0.6), Inches(0.9), Inches(11.2), Inches(1.4)
    elif t.startswith('SUBTITLE'):
        ph.left, ph.top, ph.width, ph.height = Inches(0.6), Inches(2.3), Inches(11.2), Inches(0.8)
        for p in ph.text_frame.paragraphs:
            for r in p.runs: r.font.size = Pt(24)
pic(s, D + 'c-clock-nixie.jpg', 1.17, 3.3, w=11.0)
tb(s, 0.6, 6.55, 12.1, 0.5, [('github.com/scottdube/hot-info-orbs', 16, False, GREY)], align=PP_ALIGN.CENTER)
notes(s, "Hold here for questions. The orb stays on the table running; come and press the buttons.")

prs.save(out)
print('saved', out)
