#!/usr/bin/env python3
"""BOM for the XIAO carrier board -> bom/hot_info_orbs_xiao_bom.xlsx + .csv

  python3 -m venv /tmp/bomvenv && /tmp/bomvenv/bin/pip install openpyxl
  /tmp/bomvenv/bin/python tools/make_bom.py

Schematic parts (refs, values, footprints) are from
`kicad-cli sch export bom`; everything else (sockets, headers, lenses, screws,
the board) is not in the schematic and is listed by hand. Every line has
Link 1 filled. Buy links go to Amazon (Link 1) and AliExpress
(Link 2), Scott's preference 2026-10-17; the XIAO keeps Seeed as Link 1 (plain S3 primary, Plus in Notes),
the board keeps the fab houses. "Mfr part no." still names the reference
part the footprint was drawn for, so a marketplace part can be checked
against it; where fit is not shown, Notes says so.
Prices are the Link 1 listing price seen live on 2026-10-17 (PRICES below),
before shipping and tax. "Build cost" is what one build costs out of pocket:
whole packs, so 5 displays from a 3-pack buy two packs. The per-board share
under it pro-rates each pack instead. The PCB is not priced: it is a JLCPCB
quote that depends on quantity and colors, and an unverified number would
need the +40% estimate mark.
"""
import csv
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HERE = Path(__file__).resolve().parent.parent
OUT = HERE / "bom"
OUT.mkdir(exist_ok=True)

COLS = ["Item", "Refs", "Qty / board", "Part", "Spec", "Footprint", "Mfr part no.",
        "Link 1", "Link 2", "Notes"]
AMZ = "https://www.amazon.com/dp/"
ALI = "https://www.aliexpress.us/item/"
ROWS = [
    ["U2", 1, "Seeed Studio XIAO ESP32-S3", "ESP32-S3R8, 8 MB flash, 8 MB PSRAM; SKU 113991114",
     "XIAO-ESP32-S3-SOCKET", "113991114",
     "https://www.seeedstudio.com/XIAO-ESP32S3-p-5627.html",
     AMZ + "B0BYSB66S5",
     "Primary. Seeed sells it with pin headers Un-Soldered or Pre-Soldered; Pre-Soldered makes the "
     "male header line below unnecessary. ALTERNATE if more memory is wanted: XIAO ESP32-S3 Plus "
     "(16 MB flash, SKU 102010671) " + "https://www.seeedstudio.com/Seeed-Studio-XIAO-ESP32S3-Plus-p-6361.html"
     " / " + AMZ + "B0H6XJGN3Z / " + ALI + "3256808599305867.html - ships WITHOUT headers. "
     "Same footprint and the same 8 MB firmware binary on both. Mounted on the back, on sockets."],
    ["DS1-DS5", 5, "1.28in round TFT, GC9A01, 240x240, 7-pin SPI",
     "7-pin 2.54 mm header, order VCC GND SCL SDA DC CS RST",
     "_Displays:GC9A01_SOCKET", "generic",
     AMZ + "B0DJ13VW73", ALI + "3256813175276825.html",
     "Check the listing's pin order before buying; the footprint is pad 1 = VCC (+5V) .. pad 7 = RST. "
     "The module's own male header plugs into the female socket below. Lit area ~32.5 mm dia "
     "(docs/DISPLAY-EDGE-TEST.md)."],
    ["LENS1-LENS5 (glass)", 5, "Display lens: mineral watch crystal, double dome", "35.5 mm dia, 2.0 mm thick",
     "enclosure part, no footprint", "generic",
     ALI + "3256806569387664.html", "",
     "OPTIONAL, not in the build total (priced as an add-on below it). Choose the Diameter 35.5mm "
     "option on the listing (Scott, 2026-10-17); clears the ~32.5 mm lit area. No Amazon source found yet."],
    ["LENS1-LENS5 (acrylic)", 5, "Display lens: laser-cut clear acrylic disc", "35.5 mm dia, thickness TBD",
     "enclosure part, no footprint", "cut in-house",
     "in-house (laser)", "",
     "OPTIONAL alternative to the glass lens, price TBD (Scott, 2026-10-17). Flat, so no dome. Same "
     "35.5 mm diameter as the glass; the bezel has to allow for whichever thickness is chosen."],
    ["J (displays)", 5, "Female socket 1x07, 2.54 mm, straight", "8.5 mm body, through hole",
     "PinSocket_1x07_P2.54mm_Vertical (in GC9A01_SOCKET)", "Sullins PPTC071LFBN-RC (reference)",
     AMZ + "B0GM7T6XWP", ALI + "3256812879331298.html",
     "Not a schematic part: modeled inside the display footprint. 8.5 mm is the lift the 3D assumes; "
     "the marketplace sockets' body height is NOT verified. On AliExpress choose the 7P option."],
    ["J (XIAO)", 2, "Female socket 1x07, 2.54 mm, straight", "8.5 mm body, through hole",
     "in XIAO-ESP32-S3-SOCKET", "Sullins PPTC071LFBN-RC (reference)",
     AMZ + "B0GM7T6XWP", ALI + "3256812879331298.html",
     "Same part as the display sockets; 7 per board in all. Build cost 0: they come out of the line above's 60-pack."],
    ["P (XIAO)", 2, "Male header 1x07, 2.54 mm", "cut from a 1x40 breakaway strip (1 strip = 5 sets)",
     "in XIAO-ESP32-S3-SOCKET", "Sullins PRPC040SAAN-RC (reference)",
     AMZ + "B0FRZW75VS", ALI + "3256806779943089.html",
     "Soldered to the XIAO, plugs into the XIAO sockets. Skip if the XIAO was bought Pre-Soldered."],
    ["SW1-SW3", 3, "Tactile switch, right angle, 6x6 mm", "SPST-NO, through hole, side actuated",
     "SW_Tactile_SPST_Angled_PTS645Vx39-2LFS", "C&K PTS645VL39-2 LFS (footprint part)",
     AMZ + "B008DS13P2", ALI + "3256810140983087.html",
     "FIT NOT VERIFIED: the footprint is the C&K part (7.0 mm mounting-pin pitch); the uxcell "
     "6x6x8 on Amazon lists 6.5 mm. Check one against the board before buying a lot. "
     "Exact part: https://lcsc.com/product-detail/Tactile-Switches_C-K-PTS645VL392LFS_C285523.html"],
    ["C1", 1, "Electrolytic capacitor 100 uF 16 V, +5V bulk", "100 uF 16 V, 6.3 mm dia, 2.5 mm pitch",
     "CP_Radial_D6.3mm_P2.50mm", "Panasonic EEU-FC1C101H (reference)",
     AMZ + "B0992PZC82", ALI + "3256812824344961.html",
     "Value set by Scott 2026-10-17. Both listings state 6.3 mm dia (Amazon 6.3x11, AliExpress "
     "Nichicon 6.3x7); neither states pitch, and 6.3 mm radials are 2.5 mm by convention. Mounted on "
     "the back: the 11 mm Amazon part stands 4 mm taller than the AliExpress one, so check enclosure clearance."],
    ["C2", 1, "Ceramic capacitor 1 uF (105)", "50 V, radial, 5.08 mm lead spacing",
     "C_Disc_D5.0mm_W2.5mm_P5.00mm", "KEMET C330C105K5R5TA7301 (reference)",
     AMZ + "B08BF91KY3", ALI + "3256811560978404.html",
     "Display RST hold-off with R1. Both listings are 5.08 mm pitch."],
    ["R1", 1, "Resistor 10k", "1/4 W 5 % carbon film, axial, 10.16 mm pitch",
     "R_Axial_DIN0207_L6.3mm_D2.5mm_P10.16mm_Horizontal", "Stackpole CF14JT10K0 (reference)",
     AMZ + "B0CL9KXXZB", ALI + "3256807898146236.html",
     "Display RST pull-up to 3V3."],
    ["H1-H7", 7, "M3 screw (mounting)", "M3-0.5, socket head; 3.2 mm holes",
     "Edge.Cuts circles, no footprint", "M3 x 6 DIN 912 class 12.9 (length depends on the enclosure)",
     AMZ + "B0DJQH3WSX", ALI + "3256805704775238.html",
     "Standoffs/length set by the enclosure, not chosen yet."],
    ["PCB", 1, "Bare board", "2 layer, 1.6 mm, 198.82 x 47.78 mm, WHITE solder mask, BLACK silkscreen",
     "hot_info_orbs_xiao.kicad_pcb", "-",
     "https://jlcpcb.com/quote/pcbOrderFaq/PCB%20Color", "https://www.pcbway.com/",
     "Order the colors explicitly: the stackup colors only drive the 3D/STEP, not the gerbers."],
]
assert all(len(r) == len(COLS) - 1 for r in ROWS)
assert all(r[6] for r in ROWS), "every line needs Link 1"

# Lines a builder can leave out: kept off the two build totals, priced as add-ons under them.
OPTIONAL = {"LENS1-LENS5 (glass)", "LENS1-LENS5 (acrylic)"}
TBD = "TBD"
# Refs -> (Link 1 price USD, pack qty in BOM units, build-cost override)
PRICES = {
    "U2": (7.49, 1, None),             # Seeed, headers un-soldered
    "DS1-DS5": (17.99, 3, None),       # MTDELE 3-pack
    "LENS1-LENS5 (glass)": (3.08, 1, None),  # AliExpress list price, not the new-shopper price
    "LENS1-LENS5 (acrylic)": (TBD, 1, None),
    "J (displays)": (6.99, 60, None),  # MTCELL 60 pcs
    "J (XIAO)": (6.99, 60, 0),         # out of the line above's 60-pack
    "P (XIAO)": (6.29, 100, None),     # 20 strips x 40 pins = 100 1x7 pieces
    "SW1-SW3": (7.39, 20, None),
    "C1": (6.60, 10, None),            # iProtool 6.3x11; the 50-pack gave no size
    "C2": (6.99, 120, None),
    "R1": (4.99, 100, None),
    "H1-H7": (7.49, 120, None),
}
PLUS_PRICE = 7.90                      # XIAO ESP32-S3 Plus on Seeed, seen 2026-10-17
COLS = COLS[:9] + ["Price (Link 1)", "Pack qty", "Build cost"] + COLS[9:]
OUT_ROWS = []
for r in ROWS:
    price, pack, override = PRICES.get(r[0], ("", "", None))
    OUT_ROWS.append(r[:8] + [price, pack, override] + r[8:])
NOTE = ("Required lines only; the optional lenses are below. Link 1 prices seen 2026-10-17, "
        "before shipping and tax. PCB not included: get a JLCPCB quote.")


def priced(r):
    return isinstance(r[8], (int, float))


def counted(r):
    return priced(r) and r[0] not in OPTIONAL


def build_cost(qty, price, pack, override):
    if not isinstance(price, (int, float)):
        return price
    if override is not None:
        return override
    return -(-qty // pack) * price


def share(r):
    return r[1] / r[9] * r[8]


GLASS = next(i for i, r in enumerate(OUT_ROWS) if r[0] == "LENS1-LENS5 (glass)")
U2 = next(i for i, r in enumerate(OUT_ROWS) if r[0] == "U2")
# (label, csv value, xlsx formula, note); the formulas are filled in once rows are placed
ADDONS = [
    ("Option: glass lenses, per board", build_cost(*[OUT_ROWS[GLASS][k] for k in (1, 8, 9, 10)]),
     "=L{glass}", "Five watch crystals, one per display."),
    ("Option: acrylic lenses, per board", TBD, TBD, "Laser-cut in-house; price not set yet."),
    ("Option: S3 Plus instead of S3, per board", round(PLUS_PRICE - OUT_ROWS[U2][8], 2),
     f"={PLUS_PRICE}-J{{u2}}", f"The Plus (${PLUS_PRICE:.2f} on Seeed, seen 2026-10-17) REPLACES U2; "
     "only one XIAO goes on a board."),
]

with open(OUT / "hot_info_orbs_xiao_bom.csv", "w", newline="") as f:
    w = csv.writer(f)
    w.writerow(COLS)
    for i, r in enumerate(OUT_ROWS, 1):
        w.writerow([i] + r[:10] + [build_cost(r[1], r[8], r[9], r[10])] + r[11:])
    w.writerow([])
    w.writerow([""] * 10 + ["Total, one build",
                            round(sum(build_cost(r[1], r[8], r[9], r[10]) for r in OUT_ROWS if counted(r)), 2), NOTE])
    w.writerow([""] * 10 + ["Per-board share", round(sum(share(r) for r in OUT_ROWS if counted(r)), 2),
                            "Each pack pro-rated over the boards it covers."])
    for label, val, _, note in ADDONS:
        w.writerow([""] * 10 + [label, val, note])

wb = Workbook()
ws = wb.active
ws.title = "BOM"
ws.append(COLS)
n = len(OUT_ROWS)
for i, r in enumerate(OUT_ROWS, 1):
    x = i + 1
    if not priced(r):
        bc = r[8]
    elif r[10] is not None:
        bc = r[10]
    else:
        bc = f"=ROUNDUP(C{x}/K{x},0)*J{x}"
    ws.append([i] + r[:10] + [bc] + r[11:])
ws.append([])
tot = n + 3
ws.cell(tot, 11, "Total, one build")
ws.cell(tot, 12, "=" + "+".join(f"L{i + 2}" for i, r in enumerate(OUT_ROWS) if counted(r)))
ws.cell(tot, 13, NOTE)
terms = "+".join(f"C{i + 2}/K{i + 2}*J{i + 2}" for i, r in enumerate(OUT_ROWS) if counted(r))
ws.cell(tot + 1, 11, "Per-board share")
ws.cell(tot + 1, 12, "=" + terms)
ws.cell(tot + 1, 13, "Each pack pro-rated over the boards it covers.")
for k, (label, _, formula, note) in enumerate(ADDONS):
    ws.cell(tot + 2 + k, 11, label)
    ws.cell(tot + 2 + k, 12, formula.format(glass=GLASS + 2, u2=U2 + 2))
    ws.cell(tot + 2 + k, 13, note)
# Navy header, banded body, flagged lines yellow, totals block shaded. The flag also stays
# in words (NOT VERIFIED ...) so it never rests on color alone.
NAVY, BAND, FLAG, TOTAL, RULE = "1F3A5F", "F2F5F9", "FFF2CC", "DCE6F1", "D0D7DE"
thin = Side(style="thin", color=RULE)
box = Border(left=thin, right=thin, top=thin, bottom=thin)
CENTER = (1, 3, 11)                       # Item, Qty / board, Pack qty
VENDORS = {"amazon": "Amazon", "aliexpress": "AliExpress", "seeedstudio": "Seeed",
           "jlcpcb": "JLCPCB", "pcbway": "PCBWay"}


def vendor(url):
    if not url.startswith("http"):
        return url
    host = url.split("/")[2]
    return next((v for k, v in VENDORS.items() if k in host), host)


for c in ws[1]:
    c.font = Font(bold=True, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor=NAVY)
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    c.border = box
ws.row_dimensions[1].height = 30
widths = [6, 13, 10, 34, 42, 34, 30, 14, 14, 12, 9, 12, 60]
for j, wd in enumerate(widths, 1):
    ws.column_dimensions[get_column_letter(j)].width = wd
for row in ws.iter_rows(min_row=2, max_row=n + 1):
    flagged = any(k in str(row[4].value) + str(row[12].value)
                  for k in ("NOT IN SCHEMATIC", "NOT CHOSEN", "NOT VERIFIED"))
    fill = FLAG if flagged else (BAND if row[0].row % 2 else None)
    for c in row:
        c.alignment = Alignment(wrap_text=True, vertical="top",
                                horizontal="center" if c.column in CENTER else None)
        c.border = box
        if fill:
            c.fill = PatternFill("solid", fgColor=fill)
    row[1].font = Font(bold=True)          # Refs
    for c in (row[7], row[8]):            # Link 1, Link 2: vendor name, linked
        if str(c.value).startswith("http"):
            c.hyperlink, c.value = c.value, vendor(c.value)
            c.font = Font(color="0563C1", underline="single")
    for c in (row[9], row[11]):
        c.number_format = '"$"#,##0.00'
        c.alignment = Alignment(wrap_text=True, vertical="top", horizontal="right")  # TBD lines up too
for x in range(tot, tot + 2 + len(ADDONS)):
    for j in (11, 12, 13):
        c = ws.cell(x, j)
        c.fill = PatternFill("solid", fgColor=TOTAL)
        c.border = box
        c.alignment = Alignment(wrap_text=True, vertical="top")
    main = x < tot + 2                    # the two totals; add-ons below stay plain
    ws.cell(x, 11).font = Font(bold=main, italic=not main)
    ws.cell(x, 12).font = Font(bold=main, size=12 if main else 11)
    ws.cell(x, 12).number_format = '"$"#,##0.00'
ws.sheet_properties.tabColor = NAVY
ws.freeze_panes = "C2"
ws.auto_filter.ref = f"A1:{get_column_letter(len(COLS))}{n + 1}"
wb.save(OUT / "hot_info_orbs_xiao_bom.xlsx")
print("wrote", len(ROWS), "lines to", OUT)
